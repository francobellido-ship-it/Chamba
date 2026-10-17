"""Numbered certificate annexes must be complete, ordered and validated."""

from io import BytesIO

from openpyxl import Workbook
import pymupdf
import pytest

from app.documents import prepare_document
from app.excel import _prepare_workbook
from app.excel_sheets import certificate_layouts
from app.pdf import CertificateError


def numbered_workbook(*, missing=False, duplicate=False, printed_error=False):
    book = Workbook()
    book.remove(book.active)
    for title, number in [("Certificado", 1), ("CERTIFICADO ANEXO TRES", 3), ("CERTIFICADO2", 2)]:
        if missing and number == 2:
            continue
        sheet = book.create_sheet(title)
        sheet["A1"] = "QCP-2026-0001"
        sheet["A3"] = f"Página {3 if duplicate and number == 2 else number} de 3"
        sheet["A5"] = f"DATOS FICTICIOS DE LA PAGINA {number}"
        sheet["A9"] = "17.5"
        if number == 1:
            sheet["D20"] = "Escanee este QR"
        if printed_error and number == 2:
            sheet["A9"] = "#REF!"
        sheet.print_area = "A1:L32"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = "portrait"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
    other = book.create_sheet("Registro")
    other["A1"] = "NO DEBE IMPRIMIRSE"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_numbered_annexes_are_sorted_and_the_qr_is_only_on_the_first_page():
    data = numbered_workbook()
    assert [layout.sheet_name for layout in certificate_layouts(data)] == ["Certificado", "CERTIFICADO2", "CERTIFICADO ANEXO TRES"]
    result = prepare_document(data, "xlsx", "", "https://drive.google.com/file/d/FICTICIO/view", certificate_type="no_acreditado")
    assert result.stamp_pages == [1] and result.background_pages == [1, 2, 3]
    with pymupdf.open(stream=result.data, filetype="pdf") as doc:
        assert len(doc) == 3
        for index, page in enumerate(doc):
            assert f"Página {index + 1} de 3" in page.get_text()
            assert "NO DEBE IMPRIMIRSE" not in page.get_text()
            assert len(page.get_links()) == (1 if index == 0 else 0)


@pytest.mark.parametrize("options,match", [({"missing": True}, "faltan.*2"), ({"duplicate": True}, "duplicados")])
def test_missing_or_duplicate_numbered_annexes_do_not_produce_a_partial_certificate(options, match):
    with pytest.raises(CertificateError, match=match):
        certificate_layouts(numbered_workbook(**options))


def test_printed_annex_errors_are_not_ignored():
    data = numbered_workbook(printed_error=True)
    annex = certificate_layouts(data)[1]
    with pytest.raises(CertificateError, match=r"CERTIFICADO2: A9.*#REF!"):
        _prepare_workbook(data, False, layout=annex)
