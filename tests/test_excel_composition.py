"""Numbered synthetic certificates must keep tables together without losing text."""

from io import BytesIO
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.worksheet.pagebreak import Break
import pymupdf
import pytest

from app.backgrounds import load_letterhead
from app.excel import NS, _prepare_workbook, convert_excel
from app.pdf import CertificateError
from app.excel_composition import validate_table_text, validate_page_labels
from openpyxl.drawing.image import Image as ExcelImage
from test_excel import image_bytes


def numbered_certificate(*, footer_labels=False, missing=False, conflicting_break=False, merged_across=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Certificado'
    for col in 'ABCDEFGHIJ':
        sheet.column_dimensions[col].width = 7.5
    for number in range(1, 171):
        sheet.row_dimensions[number].height = 15
    sheet['A5'] = 'FICTICIO - SIN VALIDEZ'
    sheet.merge_cells('A5:G5')
    sheet['H5'] = 'QCP-0000-0001'
    sheet.merge_cells('H5:J5')
    for number, row in [(1, 9), (2, 65), (3, 118)]:
        if missing and number == 2:
            continue
        label_row = row + (40 if footer_labels else 0)
        sheet.merge_cells(f'H{label_row}:J{label_row}')
        sheet[f'H{label_row}'] = f'Pagina {number} de 3'
    sheet['A11'] = 'DATOS FICTICIOS DEL EQUIPO'
    sheet['A20'] = 'ANTES DEL ESPACIO VACIO'
    sheet['A28'] = 'DESPUES DEL ESPACIO VACIO'
    sheet['D51'] = 'Escanee este QR'
    sheet['A67'] = 'METODO FICTICIO'
    sheet['A85'] = 'RESULTADOS FICTICIOS'
    sheet['A110'] = 'PRUEBA IBP FICTICIA'
    sheet['A120'] = 'SIGUIENTE PRUEBA FICTICIA'
    sheet['A144'] = 'OBSERVACIONES FICTICIAS'
    sheet.merge_cells('A145:J149')
    sheet['A145'] = 'Se conservan todas las observaciones del certificado ficticio.'
    sheet['A145'].alignment = Alignment(wrap_text=True)
    sheet['A151'] = 'FIN DEL DOCUMENTO FICTICIO'
    for column, label in [('A', 'Rango'), ('C', 'Patron'), ('E', 'Equipo'), ('G', 'Error'), ('I', 'Incertidumbre')]:
        for start in [111, 121]:
            sheet[f'{column}{start}'] = label
            sheet[f'{column}{start}'].fill = PatternFill('solid', fgColor='D8D8D8')
            for offset in range(1, 4):
                sheet[f'{column}{start + offset}'] = f'{offset * 40} mmHg'
            for number in range(start, start + 4):
                cell = sheet[f'{column}{number}']
                cell.font = Font(name='Arial', size=9)
                cell.border = Border(top=Side(style='thin'), bottom=Side(style='thin'),
                                     left=Side(style='thin'), right=Side(style='thin'))
                sheet.merge_cells(start_row=number, start_column=cell.column, end_row=number, end_column=cell.column + 1)
    if conflicting_break:
        sheet.row_breaks.append(Break(id=113))
    if merged_across:
        sheet.merge_cells('A64:B66')
        sheet['A64'] = 'NO CORTAR ESTA CELDA'
    sheet.print_area = 'A1:J170'
    sheet.print_title_rows = '1:8'
    sheet.print_options.horizontalCentered = True
    sheet.print_options.verticalCentered = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.scale = 82
    result = BytesIO()
    book.save(result)
    return result.getvalue()


def prepared(data):
    result, *_ = _prepare_workbook(data, False)
    with ZipFile(BytesIO(result)) as archive:
        return ET.fromstring(archive.read('xl/workbook.xml')), ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))


def test_three_logical_pages_keep_last_table_row_and_remove_orphan_before_next_page():
    result = convert_excel(numbered_certificate(), letterhead=load_letterhead('no_acreditado'))
    assert any('Se organizaron 3 páginas' in warning for warning in result.warnings)
    with pymupdf.open(stream=result.pdf, filetype='pdf') as doc:
        assert len(doc) == 3
        for i, page in enumerate(doc):
            assert f'Pagina {i+1} de 3' in page.get_text()
            assert f'Página {i+1} de 3' in page.get_text()
            assert 'SIN VALIDEZ' in page.get_text()
        assert 'PRUEBA IBP FICTICIA' in doc[1].get_text()
        assert doc[1].get_text().count('120 mmHg') == 5
        assert 'PRUEBA IBP FICTICIA' not in doc[2].get_text()
        assert 'SIGUIENTE PRUEBA FICTICIA' in doc[2].get_text()
        assert 'Se conservan todas las observaciones' in doc[2].get_text()
        assert 'FIN DEL DOCUMENTO FICTICIO' in doc[2].get_text()
        # A row split can still be extracted; check that every complete table
        # value fits within a real rendered row below its table heading.
        for page in [doc[1], doc[2]]:
            lines = [drawing['rect'] for drawing in page.get_drawings()]
            assert page.search_for('120 mmHg')
            for box in page.search_for('120 mmHg'):
                assert any(rect.y0 - 1 <= box.y0 and rect.y1 + 1 >= box.y1 and rect.height > 5 for rect in lines)


def test_empty_runs_compact_but_merged_paragraph_and_saved_values_do_not_change():
    data = numbered_certificate()
    book, sheet = prepared(data)
    area = book.find("s:definedNames/s:definedName[@name='_xlnm.Print_Area']", NS).text
    assert '$A$1:$J$51' in area and '$A$65:$J$114' in area and '$A$118:$J$151' in area
    assert sheet.find('s:printOptions', NS).get('verticalCentered') == '0'
    rows = {int(row.get('r')): row for row in sheet.findall('s:sheetData/s:row', NS)}
    assert sum(float(rows[n].get('ht')) for n in range(21, 28)) == pytest.approx(24, abs=.001)
    assert all(float(rows[n].get('ht')) == 15 for n in range(145, 150))
    with ZipFile(BytesIO(data)) as archive:
        original = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
    def values(root):
        return {cell.get('r'): ET.tostring(cell.find('s:is', NS)) if cell.find('s:is', NS) is not None else None
                for cell in root.findall('s:sheetData/s:row/s:c', NS)}
    assert values(sheet) == values(original)
    assert [m.get('ref') for m in sheet.findall('s:mergeCells/s:mergeCell', NS)] == [m.get('ref') for m in original.findall('s:mergeCells/s:mergeCell', NS)]


@pytest.mark.parametrize('options', [{'footer_labels': True}, {'conflicting_break': True}])
def test_footer_labels_and_conflicting_explicit_breaks_keep_their_layout(options):
    book, sheet = prepared(numbered_certificate(**options))
    area = book.find("s:definedNames/s:definedName[@name='_xlnm.Print_Area']", NS).text
    assert area == "'Certificado'!$A$1:$J$170"
    assert sheet.find('s:printOptions', NS).get('verticalCentered') == '1'


def test_missing_logical_page_and_merge_crossing_boundary_are_actionable():
    with pytest.raises(CertificateError, match='faltan.*2'):
        convert_excel(numbered_certificate(missing=True))
    with pytest.raises(CertificateError, match='celda combinada cruza'):
        convert_excel(numbered_certificate(merged_across=True))


def test_drawing_rows_and_original_image_bytes_survive_space_compaction():
    book = load_workbook(BytesIO(numbered_certificate()))
    book.active.add_image(ExcelImage(BytesIO(image_bytes((60, 40), 'blue'))), 'G35')
    data = BytesIO()
    book.save(data)
    original = data.getvalue()
    result, *_ = _prepare_workbook(original, False)
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(result)) as after:
        images = [name for name in before.namelist() if name.startswith('xl/media/')]
        assert images and all(before.read(name) == after.read(name) for name in images)
        sheet = ET.fromstring(after.read('xl/worksheets/sheet1.xml'))
    for number in range(35, 38):
        assert sheet.find(f"s:sheetData/s:row[@r='{number}']", NS).get('ht') == '15'


def test_single_oversized_empty_row_is_compacted_without_removing_following_text():
    book = load_workbook(BytesIO(numbered_certificate()))
    book.active.row_dimensions[29].height = 200
    book.active['A30'] = 'TEXTO DESPUES DE LA FILA MUY ALTA'
    data = BytesIO()
    book.save(data)
    _, sheet = prepared(data.getvalue())
    assert float(sheet.find("s:sheetData/s:row[@r='29']", NS).get('ht')) == 24
    assert sheet.find("s:sheetData/s:row[@r='30']/s:c/s:is/s:t", NS).text == 'TEXTO DESPUES DE LA FILA MUY ALTA'


@pytest.mark.parametrize('size,baseline', [(11, 110), (5, 109.5)])
def test_clipped_or_tiny_table_text_fails_even_when_extraction_finds_it(size, baseline):
    with pymupdf.open() as doc:
        page = doc.new_page()
        for left in [40, 150]:
            for top in [100, 112]:
                page.draw_rect(pymupdf.Rect(left, top, left + 110, top + 12))
        page.insert_text((45, baseline), 'ENCABEZADO', fontsize=size)
        assert 'ENCABEZADO' in page.get_text()
        with pytest.raises(CertificateError, match='texto de tabla recortado'):
            validate_table_text(doc)


def test_incorrect_or_clipped_numbering_is_rejected_before_emission():
    with pymupdf.open() as doc:
        doc.new_page().insert_text((40, 80), 'Pagina 1 d')
        with pytest.raises(CertificateError, match='numeración.*recortada'):
            validate_page_labels(doc, (1,), 3)


def test_small_mathematical_index_beside_a_readable_base_letter_is_preserved():
    with pymupdf.open() as doc:
        page=doc.new_page()
        for left in [40,150]:
            for top in [100,120]:
                page.draw_rect(pymupdf.Rect(left,top,left+110,top+20))
        page.insert_text((45,113),'E',fontsize=8)
        page.insert_text((50.4,113),'0',fontsize=5)
        validate_table_text(doc)
        assert 'E0' in doc[0].get_text().replace('\n','')


def test_chart_fills_do_not_create_cells_that_reject_a_complete_chart_title():
    with pymupdf.open() as doc:
        page = doc.new_page()
        # Filled chart panels have no stroked table borders. The title crosses
        # a panel boundary and must remain readable without changing the PDF.
        for left in [40, 150]:
            for top in [100, 130]:
                page.draw_rect(pymupdf.Rect(left, top, left + 110, top + 30),
                               color=None, fill=(.95, .95, .95))
        page.insert_text((125, 143), 'TITULO DEL GRAFICO', fontsize=11)
        before = page.get_text()
        validate_table_text(doc)
        assert page.get_text() == before
        assert 'TITULO DEL GRAFICO' in before


def test_background_changes_inside_a_merged_label_are_not_cell_boundaries():
    with pymupdf.open() as doc:
        page = doc.new_page()
        # A single merged label beside four data rows has two background fills,
        # but only its outer rectangle is a real cell border.
        page.draw_rect(pymupdf.Rect(40, 100, 70, 180), color=None, fill=(.8, .8, .8))
        page.draw_rect(pymupdf.Rect(40, 130, 70, 180), color=None, fill=(.75, .75, .75))
        page.draw_rect(pymupdf.Rect(40, 100, 70, 180))
        for left in [70, 180]:
            for top in [100, 120, 140, 160]:
                page.draw_rect(pymupdf.Rect(left, top, left + 110, top + 20))
        for letter, baseline in [('H', 129), ('I', 139), ('G', 149), ('H', 159)]:
            page.insert_text((50, baseline), letter, fontsize=11)
        validate_table_text(doc)
        assert page.get_text().split() == ['H', 'I', 'G', 'H']
