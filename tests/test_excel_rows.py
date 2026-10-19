"""Synthetic merged table headings: clipping, spacing and page preservation."""

from io import BytesIO
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pymupdf
import pytest

from app.excel import NS, _prepare_workbook, convert_excel
from app.excel_pages import fit_stationery_content
from test_excel import workbook_bytes


def heading_workbook(*, inferred_area=False, spacer=True, vertical_merge=False):
    book = load_workbook(BytesIO(workbook_bytes()))
    sheet = book['Certificado']
    sheet._images = []  # Isolate table sizing from drawing anchors.
    sheet['H3'] = 'QCP-2026-0001'
    sheet['D5'] = 17.5
    sheet['A40'] = 'SECOND PAGE'
    sheet['A43'] = 'TABLE SECTION'
    sheet.merge_cells('A46:C46' if not vertical_merge else 'A46:C47')
    sheet.merge_cells('D46:F46')
    for ref, text in [('A46', 'Indicación del instrumento (mm)'), ('D46', 'Incertidumbre (mm)')]:
        cell = sheet[ref]
        cell.value = text
        cell.font = Font(name='Calibri', size=12, bold=True)
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        cell.fill = PatternFill('solid', fgColor='D8D8D8')
        cell.border = Border(top=Side(style='thin'), bottom=Side(style='thin'),
                             left=Side(style='thin'), right=Side(style='thin'))
    sheet['A48'] = 'SAVED MEASUREMENT 17.5'
    if not spacer:
        sheet['A45'] = 'OCCUPIED ROW'
    if inferred_area:
        sheet.print_area = ''
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def prepared_sheet(data):
    prepared, *_ = _prepare_workbook(data, False)
    with zipfile.ZipFile(BytesIO(prepared)) as archive:
        return ET.fromstring(archive.read('xl/worksheets/sheet2.xml'))


def test_merged_heading_fits_inside_its_border_and_keeps_two_saved_pages():
    data = heading_workbook()
    prepared = prepared_sheet(data)
    source = load_workbook(BytesIO(data))['Certificado']
    rows = {int(r.get('r')): r for r in prepared.findall('s:sheetData/s:row', NS)}
    assert float(rows[46].get('ht')) > source.row_dimensions[46].height
    assert sum(float(rows[i].get('ht')) for i in range(44, 47)) == pytest.approx(42)
    with pymupdf.open(stream=convert_excel(data).pdf, filetype='pdf') as doc:
        assert len(doc) == 2
        page = doc[1]
        assert 'SECOND PAGE' in page.get_text()
        assert 'SAVED MEASUREMENT 17.5' in page.get_text()
        for text in ['Indicación del instrumento', 'Incertidumbre', '(mm)']:
            assert page.search_for(text)
        # Text extraction alone also succeeds on clipped PDFs. Require the
        # entire line box to fit inside a drawn header cell on the rendered page.
        cells = [drawing['rect'] for drawing in page.get_drawings()
                 if drawing.get('fill') and drawing['rect'].height > 10]
        for text in ['Indicación del instrumento', 'Incertidumbre', '(mm)']:
            for box in page.search_for(text):
                assert any((cell + (-1, -1, 1, 1)).contains(box) for cell in cells)


@pytest.mark.parametrize('options', [{'spacer': False}, {'vertical_merge': True}])
def test_occupied_spacers_and_vertical_merges_are_not_reflowed(options):
    sheet = prepared_sheet(heading_workbook(**options))
    row = sheet.find("s:sheetData/s:row[@r='46']", NS)
    assert float(row.get('ht')) == 14


def test_center_only_inferred_print_area_and_keep_explicit_page_settings():
    for inferred in [False, True]:
        data = heading_workbook(inferred_area=inferred)
        sheet = prepared_sheet(data)
        options = sheet.find('s:printOptions', NS)
        assert (options is not None and options.get('horizontalCentered') == '1') == inferred
        assert sheet.find('s:pageSetup', NS).get('fitToHeight') == '0'
        assert sheet.find('s:rowBreaks/s:brk', NS).get('id') == '36'


def test_fitting_to_letterhead_keeps_equal_side_margins():
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.draw_rect(pymupdf.Rect(60, 40, page.rect.width - 60, 810))
        fit_stationery_content(doc, ([], {}, False, False), added_letterhead=True)
        rect = doc[0].get_drawings()[0]['rect']
        assert rect.x0 == pytest.approx(page.rect.width - rect.x1, abs=.01)
        assert rect.y0 >= 102 and rect.y1 <= page.rect.height - 44
