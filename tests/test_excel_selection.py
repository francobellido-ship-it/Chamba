"""Fictitious alternative formats and numbered pages with an inherited code."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
import pymupdf
import pytest

from app.documents import prepare_document
from app.excel import inspect_excel
from app.excel_sheets import certificate_layouts
from app.excel_composition import validate_table_text
from app.excel_footers import FormatFooter,verify_format_footers


def numbered_formats():
    book=Workbook();book.remove(book.active)
    for number in (1,2):
        sheet=book.create_sheet(f'Certificado QCP {number}')
        sheet['A3']=f'Página {number} de 2'
        if number==1:
            sheet['A1']='QCP-2026-0001';sheet['D20']='Escanee este QR'
        else:
            sheet['A6']='PAGINA CONTINUADA SIN REPETIR EL NUMERO DEL CERTIFICADO'
            sheet['A19']='RESULTADOS';sheet['A19'].font=Font(bold=True)
            for ref,text in [('A20','Indicación del\ninstrumento\n(mm)'),('E20','Incertidumbre de\nmedición\n(mm)')]:
                end='D20'if ref=='A20'else'H20'
                sheet.merge_cells(f'{ref}:{end}')
                cell=sheet[ref];cell.value=text;cell.font=Font(name='Arial',size=8,bold=True)
                cell.alignment=Alignment(horizontal='center',vertical='center')
                cell.fill=PatternFill('solid',fgColor='D8D8D8')
                cell.border=Border(top=Side(style='thin'),bottom=Side(style='thin'),left=Side(style='thin'),right=Side(style='thin'))
            sheet['A21']='17.50';sheet['E21']='0.02'
            for ref in ['A21','E21']:
                sheet[ref].font=Font(name='Arial',size=8)
                sheet[ref].border=Border(top=Side(style='thin'),bottom=Side(style='thin'),left=Side(style='thin'),right=Side(style='thin'))
        for row in range(1,36):sheet.row_dimensions[row].height=18
        for column in 'ABCDEFGHIJKL':sheet.column_dimensions[column].width=8.3
        sheet.print_area='A1:L35';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
        sheet.page_setup.orientation='portrait';sheet.page_setup.scale=71
        sheet.sheet_properties.pageSetUpPr.fitToPage=False
    buffer=BytesIO();book.save(buffer);return buffer.getvalue()


def test_numbered_qcp_pages_share_one_pdf_and_table_labels_fit_without_clipping():
    source=numbered_formats()
    assert [p.sheet_name for p in certificate_layouts(source)]==['Certificado QCP 1','Certificado QCP 2']
    result=prepare_document(source,'xlsx','','https://example.com/test',certificate_type='no_acreditado',format_code='F-GCA-10.88',format_version='03')
    assert result.final_code=='QCP-2026-0001' and result.stamp_pages==[1]
    with pymupdf.open(stream=result.data,filetype='pdf') as doc:
        assert len(doc)==2
        validate_table_text(doc)
        verify_format_footers(doc,[FormatFooter('F-GCA-10.88','03')]*2)
        assert 'PAGINA CONTINUADA' in doc[1].get_text()
        assert '17.50' in doc[1].get_text() and '0.02' in doc[1].get_text()
        assert 'Indicación del\ninstrumento\n(mm)' in doc[1].get_text()


@pytest.mark.parametrize('current_guide',[True,False])
def test_alternative_certificates_choose_one_qcp_format_without_combining_old_codes(current_guide):
    book=Workbook();old=book.active;old.title='CERTIFICADO'
    current=book.create_sheet('CERTIFICADO QCP')
    for sheet,code in [(old,'QCP-2025-0001'),(current,'QCP-2026-0001')]:
        sheet['A1']=code;sheet['D20']='Escanee este QR';sheet.print_area='A1:L32'
    if not current_guide:current['D20']=None
    book.active=0
    buffer=BytesIO();book.save(buffer)
    selected=certificate_layouts(buffer.getvalue())
    assert len(selected)==1 and selected[0].sheet_name=='CERTIFICADO QCP'
    assert 'único formato' in inspect_excel(buffer.getvalue()).selection_reason
